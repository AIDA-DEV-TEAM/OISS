"""One writer per format. Each one puts the applied context inside the file.

The context block is generated from the same dict in every format, so a CSV, a
workbook and a PDF of the same selection disclose the same thing. Where a
format has a natural home for it -- an Excel sheet, a PDF footer -- it goes
there; where it does not, it goes into a commented header.
"""
from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping, Optional, Sequence

import matplotlib

# Server-side rendering with no display: the image is reproducible from the
# query result rather than screenshotted from a browser.
matplotlib.use("Agg")

import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import A4, landscape  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import mm  # noqa: E402
from reportlab.platypus import (  # noqa: E402
    BaseDocTemplate,
    Frame,
    Image,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from app.exports import provenance

# The PDF and PNG show a table or a chart; beyond this many rows neither is
# readable, and the tabular formats are the honest place to go for the rest.
PDF_ROW_LIMIT = 60
PNG_BAR_LIMIT = 30

Row = Mapping[str, Any]


# --------------------------------------------------------------------------
# The context block, rendered once and reused by every format
# --------------------------------------------------------------------------
def _filters_phrase(context: Mapping[str, Any]) -> str:
    filters = context.get("filters") or []
    if not filters:
        return "None"
    parts = []
    for item in filters:
        dim_label = item.get("dimension_label") or item.get("dimension") or ""
        dim_str = str(dim_label)
        if dim_str and dim_str[0].islower():
            dim_str = dim_str[0].upper() + dim_str[1:]
        values = item.get("values_display") or item.get("values") or []
        parts.append(f"{dim_str}: {', '.join(str(v) for v in values)}")
    return " / ".join(parts)


def period_phrase(context: Mapping[str, Any]) -> str:
    period = context.get("period")
    if not period:
        return "All available periods"
    start = period.get("from") or period.get("from_year")
    end = period.get("to") or period.get("to_year")
    if start and end:
        return str(start) if start == end else f"{start} to {end}"
    if start:
        return f"from {start}"
    if end:
        return f"to {end}"
    return "All available periods"


def _mix_phrase(mix: Mapping[str, Any]) -> str:
    named = {key: int(value) for key, value in (mix or {}).items() if value}
    if not named:
        return "Not recorded"
    return ", ".join(f"{key}: {value:,}" for key, value in sorted(named.items()))


def _config_phrase(config: Mapping[str, Any]) -> str:
    """One model configuration, as the model service stated it."""
    return (
        f"{config.get('model_name')}; target {config.get('target')}; "
        f"features {config.get('features')}; model version: {config.get('model_version')}; "
        f"training period: {config.get('training_period')}; "
        f"config sha256 {config.get('model_config_hash')} "
        f"(stated by {config.get('model_config_sources')}); runs {config.get('run_ids')}"
    )


def context_pairs(context: Mapping[str, Any]) -> list[tuple[str, str]]:
    """The context as ordered label/value pairs -- the one definition of it."""
    sources = context.get("source_datasets") or []
    row_count = int(context.get("row_count", 0))
    if context.get("truncated"):
        matching = int(context.get("matching_row_count") or row_count)
        row_str = f"Truncated at {row_count:,} of {matching:,} rows"
    else:
        row_str = f"{row_count:,}"

    pairs: list[tuple[str, str]] = [
        ("Panel", str(context.get("panel_title", ""))),
        ("Metric", str(context.get("metric_label") or context.get("metric") or "-")),
        ("Unit", str(context.get("unit") or "-")),
        ("Filters", _filters_phrase(context)),
        ("Period", period_phrase(context)),
        ("Rows in this file", row_str),
        ("Underlying fact rows", f"{int(context.get('underlying_row_count', 0)):,}"),
        ("Data origin", _mix_phrase(context.get("data_origin") or {})),
    ]
    grain = context.get("grain_source") or {}
    if grain:
        pairs.append(("Grain source", _mix_phrase(grain)))
    pairs.append(
        (
            "Dataset versions",
            "; ".join(str(s.get("dataset_version_id")) for s in sources) or "None",
        )
    )
    for config in context.get("model_configurations") or []:
        pairs.append(("Model configuration", _config_phrase(config)))
    for note in provenance.notes_from_context(context):
        pairs.append(("Disclosure", note))
    pairs.append(("Context", str(context.get("context_origin", "derived"))))
    pairs.append(("Generated at", str(context.get("generated_at", ""))))
    return pairs


def _caveat_lines(caveats: Sequence[Mapping[str, Any]]) -> list[str]:
    return [
        f"[{c.get('severity', 'info')}] {c.get('code')}: {c.get('message')}"
        for c in caveats
    ]


# --------------------------------------------------------------------------
# CSV -- rows under a commented context header
# --------------------------------------------------------------------------
def write_csv(
    path: Path,
    rows: Sequence[Row],
    columns: Sequence[str],
    context: Mapping[str, Any],
    caveats: Sequence[Mapping[str, Any]],
    include_context: bool = True,
    include_caveats: bool = True,
) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        if include_context:
            handle.write("# OISS export -- applied context\n")
            for label, value in context_pairs(context):
                handle.write(f"# {label}: {value}\n")
        if include_caveats and caveats:
            handle.write("#\n# Caveats\n")
            for line in _caveat_lines(caveats):
                handle.write(f"# {line}\n")
        if include_context or (include_caveats and caveats):
            handle.write("#\n")
        writer = csv.writer(handle)
        unit = str(context.get("unit") or "")
        metric = context.get("metric_label")
        writer.writerow([_header_with_unit(column, unit, metric) for column in columns])
        for row in rows:
            writer.writerow([row.get(column) for column in columns])


# --------------------------------------------------------------------------
# JSON -- the query response verbatim
# --------------------------------------------------------------------------
def write_json(
    path: Path,
    rows: Sequence[Row],
    context: Mapping[str, Any],
    caveats: Sequence[Mapping[str, Any]],
    include_context: bool = True,
    include_caveats: bool = True,
) -> None:
    body: dict[str, Any] = {"data": list(rows)}
    if include_context:
        body["applied_context"] = dict(context)
    body["caveats"] = list(caveats) if include_caveats else []
    path.write_text(json.dumps(body, indent=2, default=str), encoding="utf-8")


# --------------------------------------------------------------------------
# Excel -- Data, Context, Caveats
# --------------------------------------------------------------------------
def _typed(value: Any) -> Any:
    """Numbers as numbers and dates as dates; everything else as it came.

    An Excel export whose figures are text cannot be summed, which defeats the
    point of exporting to Excel at all.
    """
    if value is None or isinstance(value, (bool, int, float, datetime)):
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        return value


_KNOWN_COLUMN_LABELS: dict[str, str] = {
    "district": "District",
    "district_id": "District ID",
    "district_name": "District",
    "crop": "Crop",
    "crop_id": "Crop ID",
    "crop_name": "Crop",
    "crop_group": "Crop Group",
    "block": "Block",
    "block_id": "Block ID",
    "block_name": "Block",
    "month": "Month",
    "agri_year": "Agricultural Year",
    "season": "Season",
    "price_type": "Price Type",
    "price_rs_per_quintal": "Price",
    "land_use_category": "Land Use Category",
    "data_origin": "Data Origin",
    "grain_source": "Grain Source",
    "dataset_version_id": "Dataset Version ID",
    "source_file": "Source File",
    "value_status": "Value Status",
    "measure": "Measure",
    "product": "Product",
}


def _header_with_unit(column: str, unit: str, metric_label: Optional[str] = None) -> str:
    """Column headers use display names and units, per task 5 and export fix."""
    if column == "value":
        # The metric's own name: "Average price (Rs/quintal)", not "value".
        label = metric_label or "Value"
        return f"{label} ({unit})" if unit else label
    if column in _KNOWN_COLUMN_LABELS:
        base = _KNOWN_COLUMN_LABELS[column]
        return f"{base} ({unit})" if unit and column in {"price_rs_per_quintal"} else base
    if column.endswith("_id"):
        base = column[:-3]
        lbl = _KNOWN_COLUMN_LABELS.get(base, base.replace("_", " ").title())
        return f"{lbl} ID"
    return column.replace("_", " ").title()


def write_xlsx(
    path: Path,
    rows: Sequence[Row],
    columns: Sequence[str],
    context: Mapping[str, Any],
    caveats: Sequence[Mapping[str, Any]],
    include_context: bool = True,
    include_caveats: bool = True,
) -> None:
    book = Workbook()
    data = book.active
    data.title = "Data"
    unit = str(context.get("unit") or "")
    metric = context.get("metric_label")
    data.append([_header_with_unit(column, unit, metric) for column in columns])
    for cell in data[1]:
        cell.font = Font(bold=True)
    for row in rows:
        data.append([_typed(row.get(column)) for column in columns])
    for index, column in enumerate(columns, start=1):
        width = max(len(_header_with_unit(column, unit, metric)) + 2, 12)
        data.column_dimensions[get_column_letter(index)].width = min(width, 40)
    data.freeze_panes = "A2"

    if include_context:
        sheet = book.create_sheet("Context")
        sheet.append(["Field", "Value"])
        sheet["A1"].font = Font(bold=True)
        sheet["B1"].font = Font(bold=True)
        for label, value in context_pairs(context):
            sheet.append([label, value])
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 90

    if include_caveats and caveats:
        sheet = book.create_sheet("Caveats")
        sheet.append(["Severity", "Code", "Message", "Affected rows"])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for caveat in caveats:
            sheet.append(
                [
                    caveat.get("severity"),
                    caveat.get("code"),
                    caveat.get("message"),
                    _typed(caveat.get("affected_rows")),
                ]
            )
        sheet.column_dimensions["A"].width = 12
        sheet.column_dimensions["B"].width = 28
        sheet.column_dimensions["C"].width = 90
        sheet.column_dimensions["D"].width = 14

    book.save(path)


# --------------------------------------------------------------------------
# PNG -- the chart, rendered from the same result
# --------------------------------------------------------------------------
# Okabe-Ito, the same palette the dashboard uses (DESIGN.md section 2).
CHART_COLOURS = ["#0072B2", "#E69F00", "#009E73", "#CC79A7", "#56B4E9", "#D55E00"]


def write_png(
    path: Path,
    rows: Sequence[Row],
    context: Mapping[str, Any],
    caveats: Sequence[Mapping[str, Any]],
    include_context: bool = True,
    include_caveats: bool = True,
) -> None:
    dimensions = list(context.get("dimensions") or [])
    label_key = dimensions[0] if dimensions else None
    shown = list(rows)[:PNG_BAR_LIMIT]
    labels = [
        str(row.get(label_key, "")) if label_key else str(index + 1)
        for index, row in enumerate(shown)
    ]
    values = [float(row.get("value") or 0) for row in shown]

    figure, axes = plt.subplots(figsize=(11, 6), dpi=110)
    if values:
        axes.bar(labels, values, color=CHART_COLOURS[0])
        axes.set_ylabel(str(context.get("unit") or ""))
        if max(len(label) for label in labels) > 6:
            plt.setp(axes.get_xticklabels(), rotation=45, ha="right", fontsize=8)
    else:
        axes.text(0.5, 0.5, "No rows for this selection", ha="center", va="center")
        axes.set_axis_off()
    axes.set_title(str(context.get("panel_title", "")), fontsize=13, loc="left")
    axes.spines["top"].set_visible(False)
    axes.spines["right"].set_visible(False)

    # A chart image carries no rows, so its context goes in the caption -- an
    # image that leaves the system anonymous is the problem task 5 names.
    caption: list[str] = []
    if include_context:
        if context.get("truncated"):
            matching = int(context.get("matching_row_count") or context.get("row_count", 0))
            rows_phrase = f"Truncated at {int(context.get('row_count', 0)):,} of {matching:,} rows"
        else:
            rows_phrase = f"{int(context.get('row_count', 0)):,} rows"
        caption.append(
            f"Filters - {_filters_phrase(context)} | Period - {period_phrase(context)} | {rows_phrase}"
        )
        versions = "; ".join(
            str(s.get("dataset_version_id")) for s in (context.get("source_datasets") or [])
        )
        if versions:
            caption.append(f"Dataset versions - {versions}")
        caption.extend(provenance.notes_from_context(context))
    if include_caveats:
        caption.extend(_caveat_lines(caveats))
    if caption:
        figure.subplots_adjust(bottom=0.12 + 0.035 * len(caption))
        figure.text(0.02, 0.01, "\n".join(caption), fontsize=7, color="#4a4a4a", va="bottom")
    figure.savefig(path, format="png", bbox_inches="tight")
    plt.close(figure)


# --------------------------------------------------------------------------
# PDF -- a titled report
# --------------------------------------------------------------------------
def _footer(context: Mapping[str, Any]) -> str:
    versions = "; ".join(
        str(s.get("dataset_version_id")) for s in (context.get("source_datasets") or [])
    )
    parts = [f"Dataset versions: {versions or 'none'}"]
    notes = provenance.notes_from_context(context)
    if notes:
        parts.append(" / ".join(notes))
    parts.append(f"Generated {context.get('generated_at', '')}")
    return " | ".join(parts)


def write_pdf(
    path: Path,
    rows: Sequence[Row],
    columns: Sequence[str],
    context: Mapping[str, Any],
    caveats: Sequence[Mapping[str, Any]],
    include_context: bool = True,
    include_caveats: bool = True,
    chart: Optional[Path] = None,
) -> None:
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=7.5, leading=9.5)
    note = ParagraphStyle("note", parent=small, textColor=colors.HexColor("#4a4a4a"))
    footer_text = _footer(context)

    def draw_footer(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 6.5)
        canvas.setFillColor(colors.HexColor("#4a4a4a"))
        canvas.drawString(12 * mm, 10 * mm, footer_text[:240])
        canvas.drawRightString(
            document.pagesize[0] - 12 * mm, 10 * mm, f"Page {canvas.getPageNumber()}"
        )
        canvas.restoreState()

    document = BaseDocTemplate(
        str(path),
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=16 * mm,
        title=str(context.get("panel_title", "OISS export")),
    )
    frame = Frame(
        document.leftMargin,
        document.bottomMargin,
        document.width,
        document.height,
        id="body",
    )
    document.addPageTemplates(
        [PageTemplate(id="report", frames=[frame], onPage=draw_footer)]
    )

    story: list[Any] = [Paragraph(str(context.get("panel_title", "OISS export")), styles["Title"])]
    if include_context:
        for label, value in context_pairs(context):
            if label == "Panel":
                continue
            story.append(Paragraph(f"<b>{label}:</b> {value}", small))
        story.append(Spacer(1, 6))

    if chart is not None and chart.exists():
        story.append(Image(str(chart), width=document.width, height=document.width * 0.45))
        story.append(Spacer(1, 6))
    elif rows:
        shown = list(rows)[:PDF_ROW_LIMIT]
        unit = str(context.get("unit") or "")
        metric = context.get("metric_label")
        head = [Paragraph(f"<b>{_header_with_unit(c, unit, metric)}</b>", small) for c in columns]
        body = [[Paragraph(str(row.get(c, "")), small) for c in columns] for row in shown]
        table = Table([head, *body], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d4d4d4")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f2f2f2")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table)
        if len(rows) > PDF_ROW_LIMIT:
            story.append(Spacer(1, 4))
            story.append(
                Paragraph(
                    f"Showing the first {PDF_ROW_LIMIT:,} of {len(rows):,} rows. "
                    "The CSV and Excel exports of this selection carry every row.",
                    note,
                )
            )

    if include_caveats and caveats:
        story.append(Spacer(1, 8))
        story.append(Paragraph("<b>Caveats</b>", small))
        for index, line in enumerate(_caveat_lines(caveats), start=1):
            story.append(Paragraph(f"{index}. {line}", note))

    document.build(story)
